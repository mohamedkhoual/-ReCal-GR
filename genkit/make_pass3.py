"""Create the pass-3 sheets (action consistency) from files you already have.

Needs, in the same folder:
  KEY_do_not_share.csv  (from sheets.zip)
  outputs.jsonl         (from sheets.zip)
  cases.jsonl           (written to genkit/ by sample_cases.py)
Usage:  python make_pass3.py            -> writes pass3_rater1.xlsx and pass3_rater2.xlsx
"""
import json, os, sys
import pandas as pd

d = sys.argv[1] if len(sys.argv) > 1 else "."
key = pd.read_csv(os.path.join(d, "KEY_do_not_share.csv"))
outs = {(o["case_id"], o["condition"]): o["output"] for o in map(json.loads, open(os.path.join(d, "outputs.jsonl"), encoding="utf8"))}
cases = {c["case_id"]: c for c in map(json.loads, open(os.path.join(d, "cases.jsonl"), encoding="utf8"))}

def action(case, cond):
    c = cases[case]
    return {"C4": c["action"], "C5": c["action"], "C4X": c["action"], "C4M": c["mismatch_action"],
            "C4H": c["history_action"]}.get(cond)

rows = []
for _, k in key.iterrows():                     # same order as the pass-1/pass-2 sheets of rater 1
    a = action(k.case_id, k.condition)
    if a:
        rows.append({"blind_id": k.blind_id, "remediation": outs[(k.case_id, k.condition)],
                     "mandated_action": a, "action_consistency_1_5": ""})
df = pd.DataFrame(rows)

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Alignment
    xl = True
except ImportError:
    xl = False

for rater, frame in (("rater1", df), ("rater2", df.iloc[::-1])):   # rater 2 sees the reverse order
    if not xl:
        frame.to_csv(os.path.join(d, f"pass3_{rater}.csv"), index=False, encoding="utf-8-sig"); continue
    path = os.path.join(d, f"pass3_{rater}.xlsx"); frame.to_excel(path, index=False)
    wb = load_workbook(path); ws = wb.active
    dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"D2:D{ws.max_row}")
    for L, w in zip("ABCD", (12, 80, 20, 22)): ws.column_dimensions[L].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"; wb.save(path)
print(f"{len(df)} texts per rater written: pass3_rater1 and pass3_rater2")
